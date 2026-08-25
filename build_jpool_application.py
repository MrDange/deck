from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "Salmon_Wallet_JPOOL_Community_Good_Application.docx"


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="DADCE0", size="4"):
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trPr.append(tbl_header)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1155CC")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(color)
    rPr.append(underline)
    run.append(rPr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_link_line(doc, label, url):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(f"{label}: ")
    add_hyperlink(p, url, url)
    return p


def add_label_paragraph(doc, label, body):
    p = doc.add_paragraph()
    p.add_run(label).bold = True
    p.add_run(body)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.15

for style_name, size, before, after, color in (
    ("Heading 1", 20, 20, 6, "000000"),
    ("Heading 2", 16, 18, 6, "000000"),
    ("Heading 3", 14, 16, 4, "434343"),
):
    st = styles[style_name]
    st.font.name = "Arial"
    st.font.size = Pt(size)
    st.font.bold = False
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for list_name in ("List Bullet", "List Number"):
    st = styles[list_name]
    st.font.name = "Arial"
    st.font.size = Pt(11)
    st.paragraph_format.left_indent = Inches(0.5)
    st.paragraph_format.first_line_indent = Inches(-0.25)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.15

# Simple Google Docs–native title block.
title = doc.add_paragraph()
title.paragraph_format.space_before = Pt(0)
title.paragraph_format.space_after = Pt(3)
run = title.add_run("Salmon Wallet")
run.font.name = "Arial"
run.font.size = Pt(26)
run.font.bold = False
run.font.color.rgb = RGBColor(0, 0, 0)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(14)
subrun = subtitle.add_run("JPOOL Community Good Validator Application")
subrun.font.name = "Arial"
subrun.font.size = Pt(14)
subrun.font.color.rgb = RGBColor.from_string("555555")

meta = doc.add_paragraph()
meta.add_run("Validator identity: ").bold = True
meta.add_run("SA1LFXr4os2P4VKGUyRv84uFfuUYgcQkFh2uA4SmRcr")
meta.add_run("\nRegion: ").bold = True
meta.add_run("LATAM")
meta.add_run("  |  ")
meta.add_run("Project: ").bold = True
meta.add_run("Open-source Solana wallet infrastructure")

doc.add_heading("Application summary", level=1)
doc.add_paragraph(
    "Salmon Wallet is an open-source, community-owned, self-custodial wallet built exclusively for Solana. "
    "Since 2022, we have maintained user-facing wallet infrastructure through multiple market cycles, with a focus on safer signing, transaction clarity, support for emerging Solana standards, and open distribution for ecosystem protocols. "
    "Our work turns the wallet layer into reusable public infrastructure: users retain control of their assets, developers can inspect and extend the code, and Solana projects gain an additional path to reach users without depending entirely on closed platforms."
)
doc.add_paragraph(
    "Salmon is operated from Latin America, with contributors from the region, and runs the Salmon Wallet validator as an extension of that long-term commitment to Solana. "
    "JPOOL delegation would help sustain the open-source development and maintenance of this infrastructure while strengthening the geographic and operator diversity of the validator set."
)

doc.add_heading("Why Salmon qualifies as a Community Good", level=1)
for text in (
    "Open-source public infrastructure. Salmon develops a transparent, auditable wallet codebase whose improvements can be reviewed, reused, and extended by the broader ecosystem.",
    "Direct value for Solana users. The wallet provides safer self-custody and access to Solana features including Solana Pay, Token-2022, swaps, transaction simulation, protocol integrations, and Wallet Standard capabilities.",
    "Value for builders and projects. Salmon offers an open distribution surface and adapter-based integrations, reducing dependence on a small number of closed wallet channels.",
    "Long-term ecosystem commitment. The project has been actively built since 2022, received a Serum Foundation grant, and was presented at Solana Breakpoint 2022.",
    "Regional decentralization. Salmon is operated from LATAM and helps increase the region’s representation in both Solana infrastructure and user-facing products.",
):
    doc.add_paragraph(text, style="List Bullet")

doc.add_heading("Recent technical contributions", level=1)
doc.add_heading("1. Migration to @solana/kit v7", level=2)
doc.add_paragraph(
    "Salmon fully migrated its production Solana stack from the legacy @solana/web3.js v1 library to @solana/kit v7. "
    "The migration modernized signing, transfers, swaps, address handling, transaction confirmation, and dApp compatibility across the web wallet, browser extension, and mobile codebase."
)
doc.add_paragraph(
    "The work was validated through byte-for-byte golden-vector tests, cross-library verification, production builds, and real-device testing. "
    "It removed web3.js from shipped production code, strengthened type safety and RPC control, added safeguards against dependency regression, and reduced bundle size for both the web app and browser extension."
)
add_link_line(doc, "Technical documentation", "https://github.com/Salmon-HQ/salmon-wallet-frontend/blob/main/docs/SOLANA_KIT_MIGRATION.md")
add_link_line(doc, "Public announcement and technical thread", "https://x.com/salmonwallet/status/2088035108230246610?s=20")

doc.add_heading("2. OCMS v1 and safer message signing", level=2)
doc.add_paragraph(
    "Salmon implemented Solana Off-Chain Message Signing (OCMS) v1 end to end in its web wallet and browser extension. "
    "OCMS uses domain-separated signing buffers, making an off-chain message signature structurally impossible to replay as a valid Solana transaction signature."
)
doc.add_paragraph(
    "The implementation also hardens the legacy signMessage path: Salmon inspects requested bytes and refuses to sign when they deserialize as a transaction message. "
    "The approval interface proactively blocks the action rather than merely warning the user. Native Sign-In With Solana is supported and bound to the requesting origin, further reducing phishing and domain-spoofing risk."
)
doc.add_paragraph(
    "Built with Anza’s official off-chain message codec, @solana/kit, and the Wallet Standard features for signOffchainMessage and signIn, the implementation is interoperable, open-source, and covered by unit and manual end-to-end tests."
)
add_link_line(doc, "Technical documentation", "https://github.com/Salmon-HQ/salmon-wallet-frontend/blob/main/docs/OFF_CHAIN_MESSAGE_SIGNING.md")
add_link_line(doc, "Public announcement and technical thread", "https://x.com/salmonwallet/status/2086803241094754650?s=20")

doc.add_heading("Alignment with JPOOL’s Community Good criteria", level=1)
table = doc.add_table(rows=1, cols=3)
table.autofit = False
table.allow_autofit = False
widths = [Inches(1.25), Inches(1.55), Inches(3.7)]
headers = ["JPOOL area", "Salmon alignment", "Evidence"]
for idx, cell in enumerate(table.rows[0].cells):
    cell.width = widths[idx]
    cell.text = headers[idx]
    cell.paragraphs[0].runs[0].bold = True
    set_cell_margins(cell)
set_repeat_table_header(table.rows[0])

rows = [
    ("Open source", "Strong", "Public wallet codebase and technical implementation documents; development is transparent and auditable."),
    ("USERS", "Primary track", "A self-custodial wallet that improves safety and access for Solana end users."),
    ("PROJECTS", "Primary track", "An additional open distribution and integration surface for Solana protocols and applications."),
    ("DEV", "Supporting track", "Reusable implementations of modern Solana standards, documented for contributors and reviewers."),
    ("Reach and visibility", "Established", "Active since 2022; Serum grant; Breakpoint 2022 presentation; Wallet Adapter ecosystem presence; public technical releases."),
    ("Geographic impact", "LATAM", "Validator and product operated from Latin America with contributors from the region."),
]
for row_values in rows:
    cells = table.add_row().cells
    for idx, value in enumerate(row_values):
        cells[idx].width = widths[idx]
        cells[idx].text = value
        set_cell_margins(cells[idx])
set_table_borders(table)

doc.add_heading("Track record and ecosystem value", level=1)
doc.add_paragraph(
    "Salmon has remained Solana-first since its launch. The project received a grant from Serum, was showcased at Solana Breakpoint 2022, and is included in the Solana wallet-adapter ecosystem. "
    "Across successive releases, the team has expanded support for Solana Pay, Token-2022, secure signing flows, transaction simulation, swaps, domain resolution, and protocol integrations."
)
doc.add_paragraph(
    "The project’s direction goes beyond a single wallet application. Salmon is building an open access layer for Solana in which users keep custody, protocols can integrate without wallet-specific gatekeeping, and security improvements are documented and available for public review. "
    "Planned areas include native staking, validator discovery, Blinks, and further reusable integrations that connect wallet users more directly with Solana’s network and applications."
)

doc.add_heading("How JPOOL delegation would be used", level=1)
doc.add_paragraph(
    "Delegation would provide recurring, validator-aligned support for the continued development and maintenance of Salmon’s open-source wallet infrastructure. "
    "It would help fund security-focused engineering, compatibility work for new Solana standards, testing across web, extension, and mobile surfaces, and the operational resilience required to maintain both the wallet and validator over the long term."
)
doc.add_paragraph(
    "This creates a direct alignment between ecosystem contribution and network participation: the same team improving how users safely interact with Solana also operates infrastructure that helps secure and decentralize the network."
)

doc.add_heading("Application-ready project description", level=1)
doc.add_paragraph(
    "Salmon Wallet is an open-source, community-owned self-custodial wallet built exclusively for Solana. Since 2022, Salmon has developed transparent, user-first wallet infrastructure that improves signing safety, transaction clarity, support for new Solana standards, and protocol access. Recent contributions include a complete production migration from legacy @solana/web3.js to @solana/kit v7 and an end-to-end implementation of OCMS v1 that blocks transactions disguised as messages. Salmon is operated from Latin America and runs its validator as part of a long-term commitment to open infrastructure, regional representation, and a more diverse Solana network. JPOOL delegation would support continued security engineering, standard adoption, cross-platform testing, and maintenance of this public wallet infrastructure."
)

doc.add_page_break()
doc.add_heading("Validator and contact information", level=1)
details = [
    ("Validator name", "Salmon Wallet"),
    ("Validator identity", "SA1LFXr4os2P4VKGUyRv84uFfuUYgcQkFh2uA4SmRcr"),
    ("Vote account", "Sa1HXZsn2u6p2dMLZGhfxtsRw7Jo32hF15yBghWJsCz"),
    ("Client", "Jito"),
    ("Region", "LATAM"),
    ("Failover", "No"),
    ("Operator", "Nacho"),
    ("Email", "validator@salmonwallet.io"),
    ("Telegram", "@nachomileo"),
    ("X", "@salmonwallet"),
]
for label, body in details:
    add_label_paragraph(doc, f"{label}: ", body)

doc.add_heading("Supporting links", level=1)
for label, url in (
    ("Website", "https://salmonwallet.io/"),
    ("GitHub organization", "https://github.com/Salmon-HQ"),
    ("X profile", "https://x.com/salmonwallet"),
    ("Medium", "https://medium.com/@salmonwallet"),
    ("Serum Foundation grant announcement", "https://medium.com/@salmonwallet/salmon-receives-grant-from-the-serum-fund-9ef99f35b827"),
):
    add_link_line(doc, label, url)

doc.add_heading("Final submission checklist", level=1)
for item in (
    "Confirm that the validator meets JPOOL’s current baseline requirements and has an active minimum 1 SOL JSOL bond.",
    "Add current wallet usage or reach metrics if available (monthly active users, installs, connected accounts, or other verifiable adoption data).",
    "Confirm the repository’s public license and link directly to the license file in the final form.",
    "Add a Discord handle if the form requests one.",
    "Use the application-ready description above for the main narrative field and attach the technical and announcement links as impact evidence.",
):
    doc.add_paragraph(item, style="List Bullet")

doc.core_properties.title = "Salmon Wallet — JPOOL Community Good Validator Application"
doc.core_properties.subject = "Community Good validator application"
doc.core_properties.author = "Salmon Wallet"
doc.core_properties.keywords = "Solana, JPOOL, Community Good, Salmon Wallet, validator, open source"

doc.save(OUTPUT)
print(OUTPUT)
